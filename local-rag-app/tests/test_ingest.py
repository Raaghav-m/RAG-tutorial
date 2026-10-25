from pathlib import Path

import pytest
from docx import Document as WordDocument

from ingest import ingest_file


def test_ingest_file_reads_text_and_preserves_source(tmp_path: Path):
    source_path = tmp_path / "notes.md"
    source_path.write_text("Hello RAG", encoding="utf-8")

    document = ingest_file(source_path)

    assert document.text == "Hello RAG"
    assert document.source == "notes.md"
    assert document.path == str(source_path)


def test_ingest_file_rejects_unsupported_types(tmp_path: Path):
    source_path = tmp_path / "notes.csv"
    source_path.write_text("question,answer", encoding="utf-8")

    with pytest.raises(ValueError, match="unsupported file type"):
        ingest_file(source_path)


def test_ingest_file_extracts_word_paragraphs(tmp_path: Path):
    source_path = tmp_path / "notes.docx"
    word_document = WordDocument()
    word_document.add_paragraph("First paragraph")
    word_document.add_paragraph("Second paragraph")
    word_document.save(source_path)

    document = ingest_file(source_path)

    assert document.text == "First paragraph\nSecond paragraph"
    assert document.source == "notes.docx"