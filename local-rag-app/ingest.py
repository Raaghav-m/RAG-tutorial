from dataclasses import dataclass
from pathlib import Path


@dataclass
class Document:
    text: str
    source: str
    path: str


SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}


def _read_text_file(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def _read_pdf_file(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _read_docx_file(path: Path) -> str:
    from docx import Document as WordDocument

    document = WordDocument(path)
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def ingest_file(path: Path) -> Document:
    """Extract a supported local file into a common document record."""
    extension = path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"unsupported file type; expected: {supported}")

    readers = {
        ".txt": _read_text_file,
        ".md": _read_text_file,
        ".pdf": _read_pdf_file,
        ".docx": _read_docx_file,
    }

    return Document(
        text=readers[extension](path),
        source=path.name,
        path=str(path),
    )